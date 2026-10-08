"""Extract the Chettinad 15-minute Maharashtra consumption profile into average-day shapes.

Input : source-files/Chettinad_Consumption_profile.xlsx, sheet '15 min MH'
        columns Month, Day, Hour, Block(1..96), Slot(A-D), Demand (kW average for the block)
Output: src/data/seed/chettinad_profile.json
        { annualKwh, annualAvgDayKw[96], monthlyAvgDayKw{ 'MM': [96] }, monthlyKwh{ 'MM': kWh }, nRows, source }
Demand is read as kW (block average); kWh per block = kW x 0.25. Values are preserved as given (no rescaling).
"""
import json, sys
from collections import defaultdict
from openpyxl import load_workbook

src, out = sys.argv[1], sys.argv[2]
wb = load_workbook(src, data_only=True, read_only=True)
ws = wb['15 min MH']
sum_kw = defaultdict(lambda: [0.0] * 96); cnt = defaultdict(lambda: [0] * 96)
tot_kw = [0.0] * 96; tot_cnt = [0] * 96
n = 0; monthly_kwh = defaultdict(float); slot_of_block = {}; series = defaultdict(lambda: [0.0] * 96)
for row in ws.iter_rows(min_row=3, values_only=True):
    _, month, day, hour, block, slot, demand = row[:7]
    if month is None or block is None or demand is None: continue
    b = int(block) - 1; m = f"{int(month):02d}"
    sum_kw[m][b] += float(demand); cnt[m][b] += 1
    tot_kw[b] += float(demand); tot_cnt[b] += 1
    monthly_kwh[m] += float(demand) * 0.25
    slot_of_block[b] = slot; n += 1
    series[f"{int(month):02d}-{int(day):02d}"][b] = float(demand)
monthly = {m: [sum_kw[m][b] / cnt[m][b] if cnt[m][b] else 0 for b in range(96)] for m in sorted(sum_kw)}
annual = [tot_kw[b] / tot_cnt[b] if tot_cnt[b] else 0 for b in range(96)]
doc = {
    'name': 'Chettinad (Maharashtra 15-min profile)',
    'source': {'file': 'source-files/Chettinad_Consumption_profile.xlsx', 'sheet': '15 min MH', 'columns': 'Month, Day, Hour, Block, Slot, Demand', 'unitAssumption': 'Demand read as kW block-average; kWh = kW x 0.25', 'tag': 'extracted_actual'},
    'nRows': n,
    'annualKwh': sum(monthly_kwh.values()),
    'monthlyKwh': dict(sorted(monthly_kwh.items())),
    'annualAvgDayKw': annual,
    'monthlyAvgDayKw': monthly,
    'slotOfBlock': [slot_of_block.get(b) for b in range(96)],
    'daily': {k: [round(v, 2) for v in series[k]] for k in sorted(series)},
    'daysPerMonth': {m: sum(1 for k in series if k.startswith(m)) for m in sorted({k[:2] for k in series})},
}
json.dump(doc, open(out, 'w'), indent=0)
print(f"rows={n} annual={doc['annualKwh']/1e6:.3f} GWh avg={sum(annual)/96:.0f} kW peak={max(annual):.0f} kW min={min(annual):.0f} kW")
